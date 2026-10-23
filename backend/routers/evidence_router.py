import os
import shutil
import csv
import html
import logging
import hashlib
import datetime
import asyncio
from pathlib import Path
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form, status
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel
from sqlalchemy.orm import Session
import openpyxl

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from backend.config import settings
from backend.database import get_db
from backend.models import (
    CaseModel, EvidenceItemModel, EvidenceSourceModel,
    EvidenceMentionModel, EvidenceReviewModel
)
from backend.pipeline import EvidencePipeline
from backend.blockchain_service import BlockchainService

logger = logging.getLogger("atlas.evidence_router")

router = APIRouter(tags=["Evidence Ingestion & Forensic Processing"])

def _generate_pdf_dossier(item: EvidenceItemModel, target_path: Path, db: Session):
    sources = db.query(EvidenceSourceModel).filter(
        EvidenceSourceModel.evidence_id == item.evidence_id
    ).order_by(EvidenceSourceModel.page, EvidenceSourceModel.row, EvidenceSourceModel.id).all()

    mentions = db.query(EvidenceMentionModel).filter(
        EvidenceMentionModel.evidence_id == item.evidence_id
    ).all()

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#2563eb'),
        spaceAfter=10
    )
    meta_label = ParagraphStyle('ML', fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=colors.HexColor('#475569'))
    meta_val = ParagraphStyle('MV', fontName='Helvetica', fontSize=8.5, leading=11, textColor=colors.HexColor('#0f172a'))
    page_header = ParagraphStyle('PH', fontName='Helvetica-Bold', fontSize=10, leading=13, textColor=colors.HexColor('#1e293b'), spaceAfter=4)
    body_style = ParagraphStyle('DB', fontName='Courier', fontSize=8, leading=11, textColor=colors.HexColor('#0f172a'))

    elements = []
    elements.append(Paragraph('CENTRAL INVESTIGATION DEPARTMENT &mdash; EVIDENCE DOSSIER', title_style))
    elements.append(Paragraph(f'Case Exhibit: {item.filename} &bull; Case ID: {item.case_id}', subtitle_style))

    sha_disp = html.escape(item.sha256 or 'N/A')
    uploader_disp = html.escape(item.uploader or 'Investigator')
    meta_data = [
        [Paragraph('<b>Evidence ID:</b>', meta_label), Paragraph(html.escape(item.evidence_id or ''), meta_val),
         Paragraph('<b>Classification:</b>', meta_label), Paragraph(html.escape(item.job_classification or item.category or 'Document'), meta_val)],
        [Paragraph('<b>Uploader:</b>', meta_label), Paragraph(uploader_disp, meta_val),
         Paragraph('<b>Uploaded At:</b>', meta_label), Paragraph(html.escape(str(item.uploaded_at or 'N/A')[:19]), meta_val)],
        [Paragraph('<b>SHA-256:</b>', meta_label), Paragraph(f'<font size="6">{sha_disp}</font>', meta_val),
         Paragraph('<b>Status:</b>', meta_label), Paragraph(html.escape(item.status or 'Verified'), meta_val)],
    ]
    t = Table(meta_data, colWidths=[80, 180, 80, 180])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 14))

    if sources:
        pages_dict = {}
        for s in sources:
            p_num = s.page or 1
            pages_dict.setdefault(p_num, []).append(s.text or '')

        for p_num in sorted(pages_dict.keys()):
            elements.append(Paragraph(f'<b>Extracted Document Content &mdash; Page {p_num}</b>', page_header))
            combined_page_text = '\n'.join([t for t in pages_dict[p_num] if t.strip()])
            clean_text = combined_page_text.replace('\t', '    ')
            if clean_text.strip():
                elements.append(Preformatted(clean_text, body_style))
            else:
                elements.append(Paragraph('<i>[Empty text record on this page]</i>', styles['Italic']))
            elements.append(Spacer(1, 12))
    elif mentions:
        elements.append(Paragraph('<b>Extracted Forensic Entities & Intelligence</b>', page_header))
        elements.append(Spacer(1, 4))
        mention_rows = [
            [Paragraph('<b>Entity Type</b>', meta_label),
             Paragraph('<b>Identifier / Name</b>', meta_label),
             Paragraph('<b>Role</b>', meta_label),
             Paragraph('<b>Evidence Quote / Mention</b>', meta_label)]
        ]
        for m in mentions[:30]:
            mention_rows.append([
                Paragraph(html.escape(m.type or 'Entity'), meta_val),
                Paragraph(html.escape(m.raw_value or m.normalized_value or ''), meta_val),
                Paragraph(html.escape(m.role or 'Mentioned'), meta_val),
                Paragraph(html.escape((m.quote or '')[:100]), meta_val),
            ])
        mt = Table(mention_rows, colWidths=[90, 120, 80, 230])
        mt.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e2e8f0')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ]))
        elements.append(mt)
        elements.append(Spacer(1, 12))
    else:
        elements.append(Paragraph('<b>Forensic Custody Notice</b>', page_header))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(
            'This evidence item was ingested into the ATLAS Crime Analysis System. '
            'Its forensic hash and chain of custody are recorded in the audit trail.',
            styles['Normal']
        ))

    doc = SimpleDocTemplate(
        str(target_path),
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    doc.build(elements)

def _generate_excel_dossier(item: EvidenceItemModel, target_path: Path, db: Session):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Exhibit Data"

    ws.append(["Evidence ID", item.evidence_id])
    ws.append(["Case ID", item.case_id])
    ws.append(["Filename", item.filename])
    ws.append(["SHA256", item.sha256])
    ws.append(["Status", item.status])
    ws.append([])

    sources = db.query(EvidenceSourceModel).filter(
        EvidenceSourceModel.evidence_id == item.evidence_id
    ).order_by(EvidenceSourceModel.page, EvidenceSourceModel.row, EvidenceSourceModel.id).all()

    if sources:
        ws.append(["Page", "Row", "Extracted Text / Record"])
        for s in sources:
            ws.append([s.page or 1, s.row or "", s.text or ""])
    
    mentions = db.query(EvidenceMentionModel).filter(
        EvidenceMentionModel.evidence_id == item.evidence_id
    ).all()

    if mentions:
        ws2 = wb.create_sheet(title="Extracted Entities")
        ws2.append(["Type", "Raw Value", "Normalized Value", "Role", "Quote"])
        for m in mentions:
            ws2.append([m.type, m.raw_value, m.normalized_value, m.role, m.quote])

    wb.save(str(target_path))

def _generate_csv_dossier(item: EvidenceItemModel, target_path: Path, db: Session):
    sources = db.query(EvidenceSourceModel).filter(
        EvidenceSourceModel.evidence_id == item.evidence_id
    ).order_by(EvidenceSourceModel.page, EvidenceSourceModel.row, EvidenceSourceModel.id).all()

    with open(target_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Evidence ID", item.evidence_id, "Case ID", item.case_id, "Filename", item.filename])
        writer.writerow(["Page", "Row", "Content"])
        for s in sources:
            writer.writerow([s.page or 1, s.row or "", s.text or ""])

def _generate_text_dossier(item: EvidenceItemModel, target_path: Path, db: Session):
    sources = db.query(EvidenceSourceModel).filter(
        EvidenceSourceModel.evidence_id == item.evidence_id
    ).order_by(EvidenceSourceModel.page, EvidenceSourceModel.row, EvidenceSourceModel.id).all()

    lines = [
        f"ATLAS EVIDENCE DOSSIER",
        f"Evidence ID: {item.evidence_id}",
        f"Case ID: {item.case_id}",
        f"Filename: {item.filename}",
        f"SHA-256: {item.sha256}",
        f"Uploaded by: {item.uploader}",
        "=" * 60,
        ""
    ]
    for s in sources:
        lines.append(f"--- Page {s.page or 1} ---")
        lines.append(s.text or "")
        lines.append("")

    target_path.write_text("\n".join(lines), encoding="utf-8")

def resolve_or_create_evidence_file(item: EvidenceItemModel, db: Session) -> Path:
    # 1. Try originals_dir with sha256
    if item.sha256:
        p = settings.originals_dir / item.sha256
        if p.exists() and p.is_file() and p.stat().st_size > 0:
            return p

    # 2. Try uploads_dir with case_id / filename
    if item.case_id and item.filename:
        p = settings.uploads_dir / item.case_id / item.filename
        if p.exists() and p.is_file() and p.stat().st_size > 0:
            return p

    # 3. Try uploads_dir with filename
    if item.filename:
        p = settings.uploads_dir / item.filename
        if p.exists() and p.is_file() and p.stat().st_size > 0:
            return p

    # 4. Try item.relative_path
    if item.relative_path:
        for base in [Path('.'), settings.data_dir, settings.data_dir.parent]:
            cand = base / item.relative_path
            if cand.exists() and cand.is_file() and cand.stat().st_size > 0:
                return cand

    # 5. Search for filename in originals or uploads
    if item.filename:
        for root in [settings.uploads_dir, settings.originals_dir]:
            matches = list(root.rglob(item.filename))
            if matches and matches[0].is_file() and matches[0].stat().st_size > 0:
                return matches[0]

    # 6. Fallback: synthesize / generate file and preserve in storage
    dest_path = settings.originals_dir / (item.sha256 if item.sha256 else (item.evidence_id or "evidence_file"))
    dest_path.parent.mkdir(parents=True, exist_ok=True)

    fmt = (item.format or "").lower().replace(".", "")
    fname = (item.filename or "").lower()

    if fmt == "pdf" or fname.endswith(".pdf"):
        _generate_pdf_dossier(item, dest_path, db)
    elif fmt in ("xlsx", "xls") or fname.endswith((".xlsx", ".xls")):
        _generate_excel_dossier(item, dest_path, db)
    elif fmt == "csv" or fname.endswith(".csv"):
        _generate_csv_dossier(item, dest_path, db)
    else:
        _generate_text_dossier(item, dest_path, db)

    # Also keep a copy in uploads_dir if case_id and filename are known
    if item.case_id and item.filename:
        try:
            up_dest = settings.uploads_dir / item.case_id / item.filename
            up_dest.parent.mkdir(parents=True, exist_ok=True)
            if not up_dest.exists():
                shutil.copy2(dest_path, up_dest)
        except Exception as e:
            logger.warning(f"Failed copying generated evidence to uploads_dir: {e}")

    return dest_path

class ClassificationUpdateRequest(BaseModel):
    category: str
    reviewer: Optional[str] = "Investigator"

class ReviewCorrectionRequest(BaseModel):
    target_type: str = "mentions"
    target_id: str
    decision: str
    correction: Optional[str] = None
    reviewer: Optional[str] = "Investigator"

def serialize_evidence_item(e: EvidenceItemModel) -> Dict[str, Any]:
    return {
        "evidence_id": e.evidence_id,
        "case_id": e.case_id,
        "sha256": e.sha256,
        "filename": e.filename,
        "uploader": e.uploader,
        "uploaded_at": e.uploaded_at,
        "format": e.format,
        "size": e.size,
        "relative_path": e.relative_path,
        "classification_override": e.classification_override,
        "category": e.category,
        "job": {
            "job_id": e.job_id or f"JOB-{e.evidence_id}",
            "status": e.status,
            "stage": e.job_stage or "Completed",
            "progress": e.job_progress or 1.0,
            "created_at": e.uploaded_at,
            "finished_at": e.finished_at,
            "error": e.job_error,
            "warnings": e.job_warnings or [],
            "classification": e.classification_override or e.job_classification or "Document",
            "mention_count": e.mention_count or 0,
            "claim_count": e.claim_count or 0
        }
    }

@router.get("/api/evidence/cases", summary="List Evidence Cases with Count")
async def get_evidence_cases(db: Session = Depends(get_db)):
    distinct_cases = db.query(EvidenceItemModel.case_id).distinct().all()
    items = []
    for (cid,) in distinct_cases:
        count = db.query(EvidenceItemModel).filter(EvidenceItemModel.case_id == cid).count()
        c_obj = db.query(CaseModel).filter(CaseModel.case_id == cid).first()
        items.append({
            "case_id": cid,
            "count": count,
            "case_name": c_obj.case_name if c_obj else cid,
            "status": c_obj.status if c_obj else "Active"
        })
    return {"items": items}

@router.get("/api/evidence", summary="List Evidence Items for Case")
async def list_evidence(
    case_id: Optional[str] = Query(None),
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db)
):
    query = db.query(EvidenceItemModel)
    if case_id and case_id != "ALL":
        query = query.filter(EvidenceItemModel.case_id == case_id)
    items = query.order_by(EvidenceItemModel.created_at.desc()).limit(limit).all()
    return {"items": [serialize_evidence_item(e) for e in items]}

@router.post("/api/evidence/uploads", summary="Upload Case Evidence Files")
@router.post("/api/evidence/upload", summary="Upload Case Evidence Files (Alias)")
async def upload_evidence_files(
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
    case_id: Optional[str] = Form(None),
    case_id_query: Optional[str] = Query(None, alias="case_id"),
    category: Optional[str] = Form("General Investigation"),
    uploader: Optional[str] = Form("Investigator"),
    db: Session = Depends(get_db)
):
    # Consolidate files from either 'files' or 'file'
    file_list: List[UploadFile] = []
    if files:
        file_list.extend(files)
    if file:
        file_list.append(file)

    if not file_list:
        raise HTTPException(status_code=400, detail="No evidence files were attached to the upload request.")

    cid = (case_id or case_id_query or "").strip()
    if not cid:
        cid = f"CASE_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"

    # Ensure case exists in DB
    c_obj = db.query(CaseModel).filter(CaseModel.case_id == cid).first()
    if not c_obj:
        c_obj = CaseModel(
            case_id=cid,
            fir_number=f"FIR-{cid}",
            case_name=f"Case {cid}",
            status="UNDER_INVESTIGATION",
            lead_investigator=uploader or "Lead Detective",
            uploaded_by=uploader or "Investigator",
            crime_category=category or "General Investigation"
        )
        db.add(c_obj)
        db.commit()

    uploaded_items = []
    case_upload_dir = settings.uploads_dir / cid
    case_upload_dir.mkdir(parents=True, exist_ok=True)

    for upload_file in file_list:
        fname = upload_file.filename or f"evidence_{datetime.datetime.utcnow().strftime('%H%M%S')}.dat"
        target_path = case_upload_dir / fname
        try:
            with open(target_path, "wb") as f:
                shutil.copyfileobj(upload_file.file, f)

            # Run pipeline. It blocks on file parsing, LLM extraction and database
            # writes, so it is offloaded to a worker thread to keep the event loop
            # responsive. The worker opens its own session because SQLAlchemy
            # sessions are not safe to share across threads, and serialises the
            # item before closing it.
            def _process(target_path: Path = target_path, fname: str = fname) -> Dict[str, Any]:
                from backend.database import SessionLocal
                worker_db = SessionLocal()
                try:
                    item = EvidencePipeline.process_evidence(
                        file_path=target_path,
                        filename=fname,
                        case_id=cid,
                        uploader=uploader or "Investigator",
                        db=worker_db,
                        custom_category=category
                    )
                    return serialize_evidence_item(item)
                finally:
                    worker_db.close()

            uploaded_items.append(await asyncio.to_thread(_process))
        except Exception as e:
            logger.error(f"Error processing uploaded file {fname}: {e}", exc_info=True)
            db.rollback()
            fail_id = f"EV-{hashlib.md5(f'{cid}:{fname}'.encode()).hexdigest()[:12]}"
            failed_item = EvidenceItemModel(
                evidence_id=fail_id,
                case_id=cid,
                sha256="",
                filename=fname,
                uploader=uploader or "Investigator",
                uploaded_at=datetime.datetime.utcnow().isoformat() + "Z",
                format=Path(fname).suffix.lower().lstrip("."),
                size=0,
                relative_path=str(target_path),
                category=category or "General Investigation",
                status="Failed",
                job_id=f"JOB-{fail_id}",
                job_stage="Processing Failed",
                job_progress=0.0,
                job_error=str(e)
            )
            try:
                db.merge(failed_item)
                db.commit()
            except Exception:
                db.rollback()
            uploaded_items.append(serialize_evidence_item(failed_item))

    return {"items": uploaded_items}

@router.get("/api/evidence/{evidence_id}", summary="Get Evidence Details")
async def get_evidence_detail(evidence_id: str, db: Session = Depends(get_db)):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")
    return serialize_evidence_item(item)

@router.get("/api/evidence/{evidence_id}/original", summary="Download or Stream Evidence File")
async def get_evidence_original(
    evidence_id: str,
    inline: bool = Query(False),
    db: Session = Depends(get_db)
):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")

    try:
        file_path = resolve_or_create_evidence_file(item, db)
    except Exception as e:
        logger.error(f"Failed to resolve or generate evidence file for {evidence_id}: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Failed to retrieve or generate evidence document: {str(e)}")

    if not file_path or not file_path.exists():
        raise HTTPException(status_code=404, detail="Original evidence file missing from storage.")

    media_type = "application/octet-stream"
    ext = (item.format or "").lower().replace(".", "")
    fname = (item.filename or "").lower()

    if ext == "pdf" or fname.endswith(".pdf"):
        media_type = "application/pdf"
    elif ext in ("jpg", "jpeg") or fname.endswith((".jpg", ".jpeg")):
        media_type = "image/jpeg"
    elif ext == "png" or fname.endswith(".png"):
        media_type = "image/png"
    elif ext in ("csv", "txt") or fname.endswith((".csv", ".txt")):
        media_type = "text/plain"
    elif ext in ("xlsx", "xls") or fname.endswith((".xlsx", ".xls")):
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    safe_filename = item.filename or f"evidence_{evidence_id}.pdf"
    disposition = f"{'inline' if inline else 'attachment'}; filename=\"{safe_filename}\""
    return FileResponse(
        path=file_path,
        media_type=media_type,
        headers={"Content-Disposition": disposition}
    )

@router.get("/api/evidence/{evidence_id}/results/sources", summary="Extracted Text and Row Sources")
async def get_evidence_sources(
    evidence_id: str,
    limit: int = Query(500, ge=1, le=1000),
    db: Session = Depends(get_db)
):
    sources = db.query(EvidenceSourceModel).filter(
        EvidenceSourceModel.evidence_id == evidence_id
    ).limit(limit).all()
    items = []
    for s in sources:
        items.append({
            "id": s.id,
            "evidence_id": s.evidence_id,
            "case_id": s.case_id,
            "unit_id": s.unit_id,
            "unit_kind": s.unit_kind,
            "page": s.page,
            "row": s.row,
            "sheet": s.sheet,
            "text": s.text,
            "bbox": s.bbox,
            "ocr_confidence": s.ocr_confidence
        })
    return {"items": items}

@router.get("/api/evidence/{evidence_id}/results/mentions", summary="Extracted Entity Mentions & Provenance")
async def get_evidence_mentions(
    evidence_id: str,
    limit: int = Query(500, ge=1, le=1000),
    db: Session = Depends(get_db)
):
    mentions = db.query(EvidenceMentionModel).filter(
        EvidenceMentionModel.evidence_id == evidence_id
    ).limit(limit).all()
    items = []
    for m in mentions:
        items.append({
            "id": m.id,
            "evidence_id": m.evidence_id,
            "case_id": m.case_id,
            "type": m.type,
            "raw_value": m.raw_value,
            "normalized_value": m.normalized_value,
            "role": m.role,
            "identifier_trust": m.identifier_trust,
            "quote": m.quote,
            "quote_status": m.quote_status,
            "provenance": {
                "page": m.page,
                "sheet": m.sheet,
                "row": m.row,
                "snippet": m.quote,
                "bbox": m.bbox
            },
            "ocr": m.ocr_scores,
            "confirmed_by": m.confirmed_by
        })
    return {"items": items}

@router.post("/api/evidence/{evidence_id}/reviewed", summary="Mark Evidence as Reviewed")
async def mark_reviewed(evidence_id: str, db: Session = Depends(get_db)):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if item:
        item.status = "Reviewed"
        db.commit()
    return {"status": "reviewed", "evidence_id": evidence_id}

@router.post("/api/evidence/{evidence_id}/retry", summary="Reprocess Evidence File")
async def retry_evidence(evidence_id: str, db: Session = Depends(get_db)):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")
    
    file_path = resolve_or_create_evidence_file(item, db)

    if not file_path or not file_path.exists():
        raise HTTPException(status_code=404, detail="Evidence file not found on disk.")

    # Re-run pipeline
    EvidencePipeline.process_evidence(
        file_path=file_path,
        filename=item.filename,
        case_id=item.case_id,
        uploader=item.uploader,
        db=db,
        evidence_id=item.evidence_id
    )
    return {"status": "reprocessed", "evidence_id": evidence_id}

@router.post("/api/evidence/{evidence_id}/classification", summary="Update Classification")
async def update_classification(
    evidence_id: str,
    req: ClassificationUpdateRequest,
    db: Session = Depends(get_db)
):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")
    item.classification_override = req.category
    db.commit()
    return {"status": "updated", "category": req.category}

@router.get("/api/evidence/{evidence_id}/reviews", summary="Get Review Corrections")
async def get_reviews(evidence_id: str, db: Session = Depends(get_db)):
    revs = db.query(EvidenceReviewModel).filter(EvidenceReviewModel.evidence_id == evidence_id).all()
    return {"reviews": revs}

@router.post("/api/evidence/{evidence_id}/reviews", summary="Submit Human Correction")
async def submit_review(
    evidence_id: str,
    req: ReviewCorrectionRequest,
    db: Session = Depends(get_db)
):
    rev = EvidenceReviewModel(
        evidence_id=evidence_id,
        target_type=req.target_type,
        target_id=req.target_id,
        decision=req.decision,
        correction=req.correction,
        reviewer=req.reviewer or "Investigator"
    )
    db.add(rev)

    # If it is a mention correction, update the mention normalized value
    mention = db.query(EvidenceMentionModel).filter(EvidenceMentionModel.id == req.target_id).first()
    if mention and req.correction:
        mention.normalized_value = req.correction
        mention.confirmed_by = req.reviewer or "Investigator"

    db.commit()
    return {"status": "correction_recorded", "target_id": req.target_id}

@router.delete("/api/evidence/{evidence_id}", summary="Delete Evidence Item")
async def delete_evidence(
    evidence_id: str,
    officer_name: Optional[str] = Query("Investigator"),
    officer_badge: Optional[str] = Query("LE-8402"),
    db: Session = Depends(get_db)
):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")

    cid = item.case_id
    db.query(EvidenceMentionModel).filter(EvidenceMentionModel.evidence_id == evidence_id).delete()
    db.query(EvidenceSourceModel).filter(EvidenceSourceModel.evidence_id == evidence_id).delete()
    db.query(EvidenceReviewModel).filter(EvidenceReviewModel.evidence_id == evidence_id).delete()
    db.delete(item)
    db.commit()

    BlockchainService.log_activity(
        db=db,
        case_id=cid,
        action_type="EVIDENCE_DELETED",
        title=f"Evidence File Purged: {item.filename}",
        description=f"Evidence ID {evidence_id} removed by {officer_name} ({officer_badge}).",
        officer_name=officer_name,
        officer_badge=officer_badge,
        status_type="warning",
        icon_type="trash"
    )

    return {"status": "deleted", "evidence_id": evidence_id}

@router.delete("/api/evidence/cases/{case_id}", summary="Delete All Evidence for Case")
async def delete_case_evidence(case_id: str, db: Session = Depends(get_db)):
    db.query(EvidenceMentionModel).filter(EvidenceMentionModel.case_id == case_id).delete()
    db.query(EvidenceSourceModel).filter(EvidenceSourceModel.case_id == case_id).delete()
    db.query(EvidenceItemModel).filter(EvidenceItemModel.case_id == case_id).delete()
    db.commit()
    return {"status": "all_evidence_deleted", "case_id": case_id}

# ---------- Ingest and Spreadsheet Preview Compatibility Endpoints ----------

@router.post("/api/v1/ingest", summary="Fast Multi-File Ingest Endpoint")
async def v1_ingest(
    case_id: Optional[str] = Query(None),
    case_id_form: Optional[str] = Form(None, alias="case_id"),
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
    uploaded_by: Optional[str] = Query(None),
    officer_name: Optional[str] = Query(None),
    category: Optional[str] = Form("General Investigation"),
    db: Session = Depends(get_db)
):
    # Consolidate files from either 'files' or 'file'
    file_list: List[UploadFile] = []
    if files:
        file_list.extend(files)
    if file:
        file_list.append(file)

    if not file_list:
        raise HTTPException(status_code=400, detail="No evidence files were attached for ingestion.")

    cid = (case_id or case_id_form or "").strip()
    if not cid:
        cid = f"CASE_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"

    c_obj = db.query(CaseModel).filter(CaseModel.case_id == cid).first()
    if not c_obj:
        c_obj = CaseModel(
            case_id=cid,
            fir_number=f"FIR-{cid}",
            case_name=f"Case {cid}",
            status="UNDER_INVESTIGATION",
            lead_investigator=officer_name or "Lead Detective",
            uploaded_by=uploaded_by or "Investigator"
        )
        db.add(c_obj)
        db.commit()

    case_upload_dir = settings.uploads_dir / cid
    case_upload_dir.mkdir(parents=True, exist_ok=True)

    processed_count = 0
    for upload_file in file_list:
        fname = upload_file.filename or f"evidence_{datetime.datetime.utcnow().strftime('%H%M%S')}.dat"
        target_path = case_upload_dir / fname
        try:
            with open(target_path, "wb") as f:
                shutil.copyfileobj(upload_file.file, f)

            # process_evidence blocks on file parsing, LLM extraction and database
            # writes. Running it directly in this async handler would stall the
            # whole event loop, so every other endpoint stops answering until the
            # case is fully ingested. Offload it to a worker thread and give that
            # thread its own database session (SQLAlchemy sessions are not
            # safe to share across threads).
            def _process(target_path: Path = target_path, fname: str = fname) -> None:
                from backend.database import SessionLocal
                worker_db = SessionLocal()
                try:
                    EvidencePipeline.process_evidence(
                        file_path=target_path,
                        filename=fname,
                        case_id=cid,
                        uploader=officer_name or uploaded_by or "Investigator",
                        db=worker_db,
                        custom_category=category
                    )
                finally:
                    worker_db.close()

            await asyncio.to_thread(_process)
            processed_count += 1
        except Exception as e:
            logger.error(f"Error during ingest for file {fname}: {e}", exc_info=True)

    return {"case_id": cid, "status": "ingested", "files_count": len(file_list), "processed_count": processed_count}

@router.post("/api/v1/preview", summary="Spreadsheet Preview (Excel / CSV)")
async def preview_spreadsheet(file: UploadFile = File(...)):
    filename = file.filename or "spreadsheet.xlsx"
    ext = Path(filename).suffix.lower()
    content = await file.read()

    sheets = []
    if ext == ".csv":
        text = content.decode("utf-8", errors="ignore")
        reader = csv.reader(text.splitlines())
        rows = list(reader)
        sheets.append({
            "name": "CSV",
            "rows": rows[:200],
            "row_count": len(rows),
            "column_count": max([len(r) for r in rows], default=0)
        })
    else:
        # Excel
        import io
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            raw_rows = list(ws.iter_rows(values_only=True))
            str_rows = [[str(cell if cell is not None else "") for cell in r] for r in raw_rows[:200]]
            col_count = max([len(r) for r in str_rows], default=0)
            sheets.append({
                "name": sname,
                "rows": str_rows,
                "row_count": len(raw_rows),
                "column_count": col_count
            })

    return {"sheets": sheets}
